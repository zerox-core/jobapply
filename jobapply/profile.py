"""简历解析：本地文件（docx/pdf/md/txt）→ 结构化 profile。"""
from __future__ import annotations

import os

from .llm import chat_json

PROFILE_PROMPT = """你是一名简历信息抽取助手。把简历原文抽取为结构化 JSON，字段如下：
{
  "basic": {"name":"","gender":"","birth":"","phone":"","email":"","political":"","hukou":"","address":"","links":""},
  "education": [{"school":"","degree":"","major":"","start":"","end":"","gpa":"","highlights":""}],
  "internships": [{"company":"","role":"","start":"","end":"","description":"","details":[{"name":"","description":""}]}],
  "projects": [{"name":"","role":"","start":"","end":"","description":"","tech":"","links":"","details":[{"name":"","description":""}]}],
  "activities": [{"name":"","start":"","end":"","description":""}],
  "awards": [""],
  "skills": [""],
  "certs": [""],
  "self_eval": "",
  "expected": {"city":"","position":"","salary":""}
}
规则：没有的字段留空字符串或空数组；日期统一 YYYY-MM 或 YYYY-MM-DD；只输出 JSON，不要输出任何解释。
内容保全是第一原则，禁止压缩、合并或省略：
- 每段经历的 description 保留原文全部事实、数据与细节，可以大段原文照抄，按原文句子组织，不要概括成一两句。
- 一段经历内部有多个子项目 / 子方向时，逐个拆进 details 数组（name 写子项目名，description 写该子项目完整描述）。
- 校园经历 / 社团职务 / 社会实践 / 竞赛放入 activities；确属奖项荣誉的放 awards。
- self_eval 原文全量照抄。
"""


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import docx
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(t for t in parts if t.strip())
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("解析 PDF 需要 pypdf：py -m pip install pypdf") from e
        reader = PdfReader(path)
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    if ext in (".md", ".txt"):
        for enc in ("utf-8", "gbk"):
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        raise RuntimeError(f"无法识别文本编码：{path}")
    raise RuntimeError(f"不支持的简历格式：{ext}（支持 docx/pdf/md/txt）")


def structure_profile(text: str, chain) -> dict:
    messages = [
        {"role": "system", "content": PROFILE_PROMPT},
        {"role": "user", "content": text[:24000]},
    ]
    return chat_json(chain, messages)
